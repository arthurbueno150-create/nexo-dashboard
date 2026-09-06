"""Importação limitada e isolada de CSV, TSV, XLSX e XLS."""
from __future__ import annotations

import csv
import io
import multiprocessing
import zipfile
from datetime import date, datetime
from pathlib import Path

from .engine import empty

MAX_FILE = 15 * 1024 * 1024
MAX_ROWS = 50_001
MAX_COLUMNS = 200
MAX_SHEET_CELLS = 300_000
MAX_BOOK_CELLS = 500_000


def _value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _check(rows):
    if len(rows) > MAX_ROWS or any(len(row) > MAX_COLUMNS for row in rows) or sum(map(len, rows)) > MAX_SHEET_CELLS:
        raise ValueError("Limite por aba: 50.000 registros, 200 colunas e 300.000 células.")


def read_book(payload: bytes, filename: str) -> dict:
    if len(payload) > MAX_FILE:
        raise ValueError("O limite é de 15 MB por arquivo.")
    extension = Path(filename).suffix.lower()
    sheets = []
    if extension in {".csv", ".tsv"}:
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = payload.decode("cp1252")
        delimiter = "\t"
        if extension == ".csv":
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = ";" if text.count(";") >= text.count(",") else ","
        rows, cells = [], 0
        try:
            for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True):
                rows.append(row)
                cells += len(row)
                if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS or cells > MAX_SHEET_CELLS:
                    raise ValueError("A aba ultrapassa os limites de linhas, colunas ou células.")
        except csv.Error as exc:
            raise ValueError("CSV inválido: confira o delimitador e o fechamento das aspas.") from exc
        sheets = [dict(name="Dados", data=rows)]
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024 or len(archive.infolist()) > 5000:
                    raise ValueError("O Excel descompactado ultrapassa o limite de leitura.")
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            try:
                total = 0
                for sheet in workbook:
                    # Formatação distante também pode aumentar a dimensão de uma aba.
                    if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_SHEET_CELLS:
                        raise ValueError("Uma aba ultrapassa os limites. Remova linhas ou colunas distantes sem dados.")
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        rows.append([_value(v) for v in row])
                        total += len(row)
                        if total > MAX_BOOK_CELLS:
                            raise ValueError("Limite de 500.000 células por arquivo.")
                        if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS:
                            raise ValueError("A aba ultrapassa os limites de leitura.")
                    _check(rows)
                    sheets.append(dict(name=sheet.title, data=rows))
            finally:
                workbook.close()
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise ValueError("Não foi possível ler o Excel. Confira se o arquivo está íntegro e sem senha.") from exc
    elif extension == ".xls":
        import xlrd
        try:
            workbook = xlrd.open_workbook(file_contents=payload, on_demand=True)
            try:
                total = 0
                for sheet in workbook.sheets():
                    if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS or sheet.nrows * sheet.ncols > MAX_SHEET_CELLS:
                        raise ValueError("Uma aba ultrapassa os limites de leitura.")
                    total += sheet.nrows * sheet.ncols
                    if total > MAX_BOOK_CELLS:
                        raise ValueError("Limite de 500.000 células por arquivo.")
                    rows = []
                    for source in sheet.get_rows():
                        row = []
                        for cell in source:
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode).date().isoformat())
                            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                                row.append(None)
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                row.append(bool(cell.value))
                            else:
                                row.append(cell.value)
                        rows.append(row)
                    sheets.append(dict(name=sheet.name, data=rows))
            finally:
                workbook.release_resources()
        except xlrd.XLRDError as exc:
            raise ValueError("XLS inválido ou protegido por senha.") from exc
    else:
        raise ValueError("Escolha um arquivo CSV, TSV, XLSX ou XLS.")
    sheets = [s for s in sheets if any(not empty(v) for row in s["data"] for v in row)]
    if not sheets:
        raise ValueError("O arquivo não contém células preenchidas.")
    for sheet in sheets:
        _check(sheet["data"])
    if sum(len(row) for s in sheets for row in s["data"]) > MAX_BOOK_CELLS:
        raise ValueError("Limite de 500.000 células por arquivo.")
    return dict(name=Path(filename.replace("\\", "/")).name[:180], demo=None, sheets=sheets)


def _worker(connection, payload, filename):
    try:
        result = (True, read_book(payload, filename))
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else "Não foi possível ler o arquivo. Confira se ele está íntegro e sem senha."
        result = (False, message)
    try:
        connection.send(result)
    except (BrokenPipeError, OSError):
        # O servidor pode ter encerrado a conexão após o limite de tempo.
        pass
    finally:
        connection.close()


def import_book(payload: bytes, filename: str, timeout: float = 30) -> dict:
    """Interrompe o processo de leitura se ultrapassar o prazo."""
    if len(payload) > MAX_FILE:
        raise ValueError("O limite é de 15 MB por arquivo.")
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=_worker, args=(send, payload, filename), daemon=True)
    process.start()
    send.close()
    try:
        if not receive.poll(timeout):
            raise ValueError("A leitura ultrapassou 30 segundos. Tente um arquivo menor.")
        try:
            success, result = receive.recv()
        except EOFError as exc:
            raise ValueError("O processo de leitura foi interrompido.") from exc
        if not success:
            raise ValueError(result)
        return result
    finally:
        receive.close()
        process.join(timeout=1)
        if process.is_alive():
            process.terminate()
            process.join(timeout=2)
