"""Leitura incremental para SQLite. Nenhuma matriz completa atravessa processos."""
from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
import time
import zipfile
from contextlib import closing
from itertools import chain, islice
from pathlib import Path

from .engine import detect_header, empty, normalize, number_value
from .imports import _value

FORMATS = {'.csv', '.tsv', '.xlsx', '.xlsm', '.xltx', '.xltm', '.xls', '.xlsb', '.ods'}
MAX_ROWS = 2_000_000
MAX_CELLS = 20_000_000
MAX_COLUMNS = 200
MAX_CELL_TEXT = 65_536
MAX_DB = 768 * 1024 * 1024


def zip_check(path):
    with zipfile.ZipFile(path) as archive:
        items = archive.infolist()
        if len(items) > 5000 or sum(x.file_size for x in items) > 1024 * 1024 * 1024:
            raise ValueError('O conteúdo descompactado excede 1 GB ou 5.000 componentes.')
        if any('sharedstrings' in x.filename.lower() and x.file_size > 32 * 1024 * 1024 for x in items):
            raise ValueError('O índice de textos do Excel é muito grande para este servidor. Exporte a aba como CSV.')


def ods_sheets(path):
    """ODS via XML incremental, inclusive células/linhas repetidas e mescladas."""
    from defusedxml.ElementTree import iterparse
    ns = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
    office = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'
    t = '{' + ns + '}'
    with zipfile.ZipFile(path) as archive, archive.open('content.xml') as source:
        sheet = None
        table = None
        for event, element in iterparse(source, events=('start', 'end'), forbid_dtd=True):
            if event == 'start' and element.tag == t + 'table':
                sheet = element.get(t + 'name', 'Dados')
                table = element
            elif event == 'end' and element.tag == t + 'table-row' and sheet is not None:
                row = []
                for cell in element:
                    if cell.tag not in {t + 'table-cell', t + 'covered-table-cell'}:
                        continue
                    value = cell.get(office + 'value') or cell.get(office + 'date-value')
                    if value is None:
                        value = cell.get(office + 'boolean-value') or cell.get(office + 'string-value')
                    if value is None:
                        value = '\n'.join(''.join(p.itertext()) for p in cell if p.tag.endswith('}p')) or None
                    repeat = int(cell.get(t + 'number-columns-repeated', '1'))
                    if repeat < 1 or repeat + len(row) > MAX_COLUMNS:
                        # LibreOffice preenche o fim da linha com células vazias repetidas.
                        if value is None and repeat > 0:
                            repeat = max(0, MAX_COLUMNS + 1 - len(row))
                        else:
                            raise ValueError('Limite de 200 colunas por aba.')
                    row.extend([value] * repeat)
                while row and empty(row[-1]):
                    row.pop()
                repetitions = int(element.get(t + 'number-rows-repeated', '1'))
                if repetitions < 1:
                    raise ValueError('Repetição inválida no ODS.')
                if row:
                    if repetitions > MAX_ROWS:
                        raise ValueError('Limite de 2 milhões de registros por arquivo.')
                    for _ in range(repetitions):
                        yield sheet, row
                element.clear()
                if table is not None:
                    table.clear()
            elif event == 'end' and element.tag == t + 'table':
                sheet = None
                element.clear()


def iter_sheets(path, extension):
    if extension in {'.csv', '.tsv'}:
        csv.field_size_limit(MAX_CELL_TEXT)
        # Valida a codificação em blocos para não duplicar o arquivo em memória.
        import codecs
        decoder = codecs.getincrementaldecoder('utf-8-sig')()
        encoding = 'utf-8-sig'
        try:
            with open(path, 'rb') as source:
                for block in iter(lambda: source.read(1024 * 1024), b''):
                    decoder.decode(block)
                decoder.decode(b'', final=True)
        except UnicodeDecodeError:
            encoding = 'cp1252'
        with open(path, encoding=encoding, newline='') as source:
            sample = source.read(65536)
            source.seek(0)
            delimiter = '\t'
            if extension == '.csv':
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t').delimiter
                except csv.Error:
                    delimiter = ';' if sample.count(';') >= sample.count(',') else ','
            for row in csv.reader(source, delimiter=delimiter, strict=True):
                yield 'Dados', row
    elif extension in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        from openpyxl import load_workbook
        zip_check(path)
        with closing(load_workbook(path, read_only=True, data_only=True, keep_links=False, keep_vba=False)) as book:
            for sheet in book:
                if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 30:
                    raise ValueError('Dimensões do Excel excedem os limites. Remova formatação distante ou exporte para CSV.')
                for row in sheet.iter_rows(values_only=True):
                    yield sheet.title, [_value(v) for v in row]
    elif extension == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for index, name in enumerate(book.sheet_names()):
                sheet = book.sheet_by_index(index)
                if sheet.ncols > MAX_COLUMNS:
                    raise ValueError('Limite de 200 colunas por aba.')
                for source in sheet.get_rows():
                    yield name, [xlrd.xldate_as_datetime(c.value, book.datemode).date().isoformat()
                                 if c.ctype == xlrd.XL_CELL_DATE else None
                                 if c.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}
                                 else bool(c.value) if c.ctype == xlrd.XL_CELL_BOOLEAN else c.value for c in source]
                book.unload_sheet(index)
        finally:
            book.release_resources()
    elif extension == '.xlsb':
        from pyxlsb import open_workbook
        zip_check(path)
        with open_workbook(str(path)) as book:
            for name in book.sheets:
                with book.get_sheet(name) as sheet:
                    if sheet.dimension and sheet.dimension.w > MAX_COLUMNS:
                        raise ValueError('Limite de 200 colunas por aba.')
                    for row in sheet.rows(sparse=True):
                        if row and max(c.c for c in row) >= MAX_COLUMNS:
                            raise ValueError('Limite de 200 colunas por aba.')
                        values = [None] * (max((c.c for c in row), default=-1) + 1)
                        for cell in row:
                            values[cell.c] = _value(cell.v)
                        yield name, values
    elif extension == '.ods':
        zip_check(path)
        yield from ods_sheets(path)
    else:
        raise ValueError('Formato não suportado. Use CSV, TSV, XLSX, XLSM, XLTX, XLTM, XLS, XLSB ou ODS.')


def progress(folder, **values):
    path = Path(folder) / 'progress.json'
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(values, ensure_ascii=False), encoding='utf-8')
    temporary.replace(path)


def build_database(folder, filename, header=None):
    """Executado em processo separado. Persiste lotes pequenos e metadados."""
    folder = Path(folder)
    source = folder / ('source' + Path(filename).suffix.lower())
    db = None
    stream = None
    try:
        progress(folder, state='processing', rows=0, message='Lendo as abas…')
        db = sqlite3.connect(folder / 'data.sqlite')
        db.execute('PRAGMA cache_size=-4096')
        db.execute('PRAGMA temp_store=FILE')
        db.execute(f'PRAGMA max_page_count={MAX_DB // 4096}')
        db.executescript('CREATE TABLE rows(sheet INTEGER, n INTEGER, payload TEXT, search TEXT, PRIMARY KEY(sheet,n));'
                         'CREATE TABLE sheets(id INTEGER PRIMARY KEY, name TEXT, columns TEXT, count INTEGER, header INTEGER);')
        from itertools import groupby
        total = cells = 0
        stream = iter_sheets(source, source.suffix)
        for sid, (name, records) in enumerate(groupby(stream, key=lambda x: x[0])):
            if sid >= 50:
                raise ValueError('Limite de 50 abas por arquivo.')
            iterator = (r for _, r in records)
            sample = list(islice(iterator, 30))
            if not sample:
                continue
            if any(len(r) > MAX_COLUMNS for r in sample):
                raise ValueError('Limite de 200 colunas por aba, incluindo o cabeçalho.')
            if any(isinstance(v, str) and len(v) > MAX_CELL_TEXT for r in sample for v in r):
                raise ValueError('Uma célula excede 65.536 caracteres.')
            selected = detect_header(sample) if header is None else header
            if selected >= len(sample):
                raise ValueError('A linha de cabeçalho escolhida não existe em uma das abas.')
            columns = [str(v).strip() if not empty(v) else f'Coluna {i+1}' for i, v in enumerate(sample[selected])]
            count = 0
            for row in chain(sample[selected+1:], iterator):
                if len(row) > MAX_COLUMNS:
                    raise ValueError('Limite de 200 colunas por aba.')
                if not any(not empty(v) for v in row):
                    continue
                if any(isinstance(v, str) and len(v) > MAX_CELL_TEXT for v in row):
                    raise ValueError('Uma célula excede 65.536 caracteres. Reduza o texto ou divida a coluna.')
                while len(columns) < len(row):
                    columns.append(f'Coluna {len(columns)+1}')
                total += 1
                cells += max(len(row), len(columns))
                count += 1
                if total > MAX_ROWS or cells > MAX_CELLS:
                    raise ValueError('Limite de 2 milhões de registros ou 20 milhões de células por arquivo.')
                payload = json.dumps(row, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
                if len(payload.encode('utf-8')) > 1024 * 1024:
                    raise ValueError('Uma linha excede 1 MB de texto. Divida as colunas de texto.')
                db.execute('INSERT INTO rows VALUES (?,?,?,?)', (sid, count, payload, normalize(' '.join(str(v) for v in row if v is not None))))
                if total % 5000 == 0:
                    db.commit()
                    progress(folder, state='processing', rows=total, message=f'Importando a aba {name}…')
            db.execute('INSERT INTO sheets VALUES (?,?,?,?,?)', (sid, name, json.dumps(columns, ensure_ascii=False), count, selected+1))
            rectangular_cells = sum(s[0] * len(json.loads(s[1])) for s in db.execute('SELECT count, columns FROM sheets'))
            if rectangular_cells > MAX_CELLS:
                raise ValueError('Limite de 20 milhões de células por arquivo, incluindo lacunas nas linhas.')
        if not total:
            raise ValueError('O arquivo não contém registros após o cabeçalho.')
        db.commit()
        progress(folder, state='ready', rows=total, message='Importação concluída.')
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else (
            'Não foi possível processar o arquivo dentro dos recursos disponíveis. '
            'Confira o formato, a codificação e a senha; tente CSV ou uma planilha menor.')
        progress(folder, state='failed', rows=0, message=message)
    finally:
        if stream is not None:
            stream.close()
        if db is not None:
            db.close()
        source.unlink(missing_ok=True)
        try:
            if json.loads((folder / 'progress.json').read_text(encoding='utf-8'))['state'] != 'ready':
                (folder / 'data.sqlite').unlink(missing_ok=True)
        except (OSError, ValueError):
            pass


def worker(folder, filename, header):
    if os.name == 'posix':
        import resource
        # Mantém espaço para o servidor dentro dos 512 MB do plano gratuito.
        resource.setrlimit(resource.RLIMIT_AS, (320 * 1024 * 1024, 320 * 1024 * 1024))
    build_database(folder, filename, header)


class DecimalSum:
    def __init__(self):
        from decimal import Decimal
        self.total = Decimal(0)

    def step(self, value):
        parsed = number_value(value)
        if parsed is not None:
            self.total += parsed

    def finalize(self):
        return str(self.total)


def connect(folder):
    db = sqlite3.connect(f'{(Path(folder) / "data.sqlite").as_uri()}?mode=ro', uri=True)
    db.execute('PRAGMA cache_size=-4096')
    db.execute('PRAGMA temp_store=FILE')
    deadline = time.monotonic() + 90
    db.set_progress_handler(lambda: int(time.monotonic() > deadline), 10000)
    db.create_aggregate('decimal_sum', 1, DecimalSum)
    db.create_function('is_number', 1, lambda v: int(number_value(v) is not None))
    return db


def selection(db, args):
    sheets = [dict(id=s[0], name=s[1], columns=json.loads(s[2]), count=s[3], header=s[4]) for s in db.execute('SELECT * FROM sheets ORDER BY id')]
    def integer(key, default):
        try:
            return int(args.get(key, default))
        except (TypeError, ValueError):
            return default
    sheet = next((s for s in sheets if s['id'] == integer('sheet', sheets[0]['id'])), sheets[0])
    metric = integer('metric', -1)
    category = integer('group', -1)
    for value in (metric, category):
        if value < -1 or value >= len(sheet['columns']):
            raise ValueError('Coluna inválida.')
    where = 'sheet=?'
    params = [sheet['id']]
    search = str(args.get('search', ''))[:500]
    if search:
        where += ' AND instr(search,?)>0'
        params.append(normalize(search))
    category_value = str(args.get('category', ''))[:500]
    if category >= 0 and category_value:
        where += f' AND CAST(json_extract(payload,\'$[{category}]\') AS TEXT)=?'
        params.append(category_value)
    count = db.execute(f'SELECT count(*) FROM rows WHERE {where}', params).fetchone()[0]
    pages = max(1, (count + 49) // 50)
    page = min(max(1, integer('page', 1)), pages)
    result = dict(sheets=sheets, sheet=sheet, metric=metric, group=category, search=search,
                  category=category_value, count=count, pages=pages, page=page, total=None,
                  valid=0, average=None, where=where, params=params)
    if metric >= 0:
        expression = f'json_extract(payload,\'$[{metric}]\')'
        total, valid = db.execute(f'SELECT decimal_sum({expression}), sum(is_number({expression})) FROM rows WHERE {where}', params).fetchone()
        result.update(total=total or '0', valid=valid or 0)
        if valid:
            from decimal import Decimal
            result['average'] = str(Decimal(total) / valid)
    result['rows'] = [json.loads(r[0]) for r in db.execute(f'SELECT payload FROM rows WHERE {where} ORDER BY n LIMIT 50 OFFSET ?', params + [(page-1)*50])]
    return result


def csv_chunks(db, selected):
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + value
        return '' if value is None else value
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, delimiter=';', quoting=csv.QUOTE_ALL, lineterminator='\r\n')
    yield '\ufeff'
    writer.writerow([safe(v) for v in selected['sheet']['columns']])
    yield stream.getvalue()
    stream.seek(0)
    stream.truncate(0)
    for index, (payload,) in enumerate(db.execute(f'SELECT payload FROM rows WHERE {selected["where"]} ORDER BY n', selected['params'])):
        row = json.loads(payload)
        writer.writerow([safe(row[i] if i < len(row) else None) for i in range(len(selected['sheet']['columns']))])
        if index % 500 == 499:
            yield stream.getvalue()
            stream.seek(0)
            stream.truncate(0)
    if stream.tell():
        yield stream.getvalue()
